import uuid
from io import BytesIO
from unittest.mock import patch
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient
from modules.accounts.models import User
from modules.resumes.models import CVDocument
from .models import ExportJob
from .rendering import render_cv_pdf
from .templates import render_html
from .docx_rendering import render_cv_docx
from docx import Document

class TemplateContractTests(TestCase):
    def test_docx_keeps_content_editable_and_marks_rtl_paragraphs(self):
        content = render_cv_docx({"template_key": "professional-v1", "design_options": {"theme": "maroon", "font": "serif", "layout": "sidebar"}, "output_language": "fa", "title": "رزومه",
            "snapshot": {"personal": {"full_name": "نام کاربر", "summary": "متن قابل ویرایش"},
                "experience": [{"position": "انجنیر", "duties": "شرح وظایف"}]}, "selected_sections": []})
        self.assertTrue(content.startswith(b"PK"))
        document = Document(BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs + [paragraph for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs])
        self.assertIn("نام کاربر", text)
        self.assertIn("شرح وظایف", text)
        self.assertTrue(any(paragraph._p.pPr is not None and paragraph._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi") is not None for paragraph in document.paragraphs))
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.tables[0].rows[0].cells), 2)

    def test_template_theme_and_layout_options_render_and_reject_invalid_values(self):
        from .templates import validate_design_options
        html = render_html({"template_key": "professional-v1", "output_language": "en", "design_options": {"theme": "maroon", "font": "serif", "layout": "sidebar"},
            "snapshot": {"personal": {"full_name": "A"}, "skills": [{"skills": ["Writing"]}], "experience": [{"position": "Engineer"}]}})
        self.assertIn("layout-sidebar", html)
        self.assertIn("#4b2634", html)
        with self.assertRaises(ValueError):
            validate_design_options("ats-v1", {"theme": "green", "font": "sans", "layout": "sidebar"})

    def test_short_docx_fixtures_cover_all_templates_and_languages(self):
        for template_key in ("public-service-v1", "professional-v1", "ats-v1"):
            for language in ("en", "fa", "ps"):
                layout = "single-column" if template_key == "ats-v1" else "sidebar" if template_key == "professional-v1" else "classic"
                output = render_cv_docx({"template_key": template_key, "design_options": {"theme": "blue-teal", "font": "sans", "layout": layout},
                    "output_language": language, "snapshot": {"personal": {"full_name": "Editable Name"}, "experience": [{"position": "Engineer", "duties": "Editable duties"}]},
                    "selected_sections": []})
                doc = Document(BytesIO(output))
                paragraphs = doc.paragraphs + [paragraph for table in doc.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
                self.assertIn("Editable Name", "\n".join(p.text for p in paragraphs))
                self.assertIn("Editable duties", "\n".join(p.text for p in paragraphs))

    def test_photo_is_included_in_regular_templates_and_omitted_from_ats(self):
        import base64
        from io import BytesIO as ImageBuffer
        from PIL import Image
        image = Image.new("RGB", (24, 32), "#17324d")
        data = ImageBuffer()
        image.save(data, format="JPEG")
        encoded = base64.b64encode(data.getvalue()).decode("ascii")
        regular = render_html({"template_key": "professional-v1", "output_language": "en", "photo_data": encoded, "snapshot": {"personal": {"full_name": "A"}}})
        ats = render_html({"template_key": "ats-v1", "output_language": "en", "photo_data": encoded, "snapshot": {"personal": {"full_name": "A"}}})
        self.assertIn("data:image/jpeg;base64,", regular)
        self.assertNotIn("data:image/jpeg;base64,", ats)
        with_photo = Document(BytesIO(render_cv_docx({"template_key": "professional-v1", "output_language": "en", "photo_data": encoded, "snapshot": {"personal": {"full_name": "A"}}})))
        without_photo = Document(BytesIO(render_cv_docx({"template_key": "ats-v1", "output_language": "en", "photo_data": encoded, "snapshot": {"personal": {"full_name": "A"}}})))
        self.assertEqual(len(with_photo.inline_shapes), 1)
        self.assertEqual(len(without_photo.inline_shapes), 0)

    def test_plain_text_is_escaped_and_rtl_language_is_explicit(self):
        html = render_html({"title": "<script>bad()</script>", "output_language": "fa", "snapshot": {"personal": {"full_name": "<script>bad()</script>"}}, "selected_sections": []})
        self.assertIn("dir='rtl'", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertNotIn("<script>bad", html)

    def test_licensed_bundled_font_and_multilingual_multipage_pdf_render(self):
        long_text = "متن رزومه sample résumé " * 350
        for template_key in ("public-service-v1", "professional-v1", "ats-v1"):
            for language in ("en", "fa", "ps"):
                design = {"theme": "blue-teal", "font": "sans", "layout": "single-column" if template_key == "ats-v1" else "classic"}
                pdf = render_cv_pdf({"title": "CV", "output_language": language, "template_key": template_key, "design_options": design,
                    "selected_sections": ["personal", "experience"], "snapshot": {"personal": {"full_name": "Nasrat Nasrati", "summary": long_text},
                    "experience": [{"position": "Engineer", "employer": "Example", "duties": long_text}]}})
                self.assertTrue(pdf.startswith(b"%PDF-"))
                self.assertGreater(len(pdf), 5000)

@override_settings(MEDIA_ROOT="/tmp/kabulcv-export-tests", CACHES={"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache"}})
class ExportApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="exporter@example.com", password="StrongPass123!", is_active=True)
        self.other = User.objects.create_user(email="other-exporter@example.com", password="StrongPass123!", is_active=True)
        self.document = CVDocument.objects.create(owner=self.user, title="Public CV", snapshot={"personal": {"full_name": "Nasrat"}}, selected_sections=["personal"])
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    @patch("modules.exports.views.render_export.delay")
    def test_preview_job_is_immutable_idempotent_and_queued_without_download_credit(self, delay):
        key = str(uuid.uuid4())
        body = {"document_id": str(self.document.id), "purpose": "preview", "idempotency_key": key}
        with self.captureOnCommitCallbacks(execute=True):
            first = self.client.post("/api/v1/exports/", body, format="json")
            again = self.client.post("/api/v1/exports/", body, format="json")
        self.assertEqual(first.status_code, 202)
        self.assertEqual(again.data["id"], first.data["id"])
        job = ExportJob.objects.get(pk=first.data["id"])
        self.assertEqual(job.purpose, ExportJob.PREVIEW)
        self.assertEqual(job.document_revision, self.document.revision)
        self.assertEqual(job.render_input["snapshot"]["personal"]["full_name"], "Nasrat")
        self.assertEqual(delay.call_count, 1)
        with self.assertRaises(ValueError):
            job.render_input = {"snapshot": {}}
            job.save()

    @patch("modules.exports.views.render_export.delay")
    def test_incomplete_export_is_rejected_but_preview_can_be_requested(self, delay):
        self.document.snapshot = {}
        self.document.save(update_fields=["snapshot"])
        response = self.client.post("/api/v1/exports/", {"document_id": str(self.document.id), "purpose": "export", "idempotency_key": str(uuid.uuid4())}, format="json")
        self.assertEqual(response.status_code, 400)
        with self.captureOnCommitCallbacks(execute=True):
            preview = self.client.post("/api/v1/exports/", {"document_id": str(self.document.id), "purpose": "preview", "idempotency_key": str(uuid.uuid4())}, format="json")
        self.assertEqual(preview.status_code, 202)
        self.assertEqual(delay.call_count, 1)

    def test_free_user_cannot_request_docx_and_verified_premium_can(self):
        from datetime import timedelta
        from django.utils import timezone
        from modules.billing.models import Plan, Subscription
        body = {"document_id": str(self.document.id), "format": "docx", "template_key": "professional-v1", "purpose": "export", "idempotency_key": str(uuid.uuid4())}
        denied = self.client.post("/api/v1/exports/", body, format="json")
        self.assertEqual(denied.status_code, 403)
        now = timezone.now()
        Subscription.objects.create(owner=self.user, plan=Plan.objects.get(code="premium"), starts_at=now-timedelta(minutes=1),
            expires_at=now+timedelta(days=5), verified_at=now)
        with patch("modules.exports.views.render_export.delay"), patch("modules.exports.views.readiness", return_value={"ready": True, "missing_fields": []}):
            with self.captureOnCommitCallbacks(execute=True):
                allowed = self.client.post("/api/v1/exports/", {**body, "idempotency_key": str(uuid.uuid4())}, format="json")
        self.assertEqual(allowed.status_code, 202)
        self.assertEqual(ExportJob.objects.get(pk=allowed.data["id"]).format, ExportJob.DOCX)
        self.assertEqual(ExportJob.objects.get(pk=allowed.data["id"]).template_key, "professional-v1")

    def test_template_catalog_and_design_options_are_plan_aware_and_persisted(self):
        catalog = self.client.get("/api/v1/exports/templates/")
        self.assertEqual({item["key"] for item in catalog.data}, {"public-service-v1", "professional-v1", "ats-v1"})
        self.assertFalse(next(item for item in catalog.data if item["key"] == "professional-v1")["is_available"])
        invalid = self.client.patch(f"/api/v1/resumes/{self.document.id}/", {"expected_revision": self.document.revision,
            "request_id": str(uuid.uuid4()), "template_key": "ats-v1", "design_options": {"theme": "green", "font": "sans", "layout": "sidebar"}}, format="json")
        self.assertEqual(invalid.status_code, 400)
        updated = self.client.patch(f"/api/v1/resumes/{self.document.id}/", {"expected_revision": self.document.revision,
            "request_id": str(uuid.uuid4()), "template_key": "professional-v1", "design_options": {"theme": "maroon", "font": "serif", "layout": "sidebar"}}, format="json")
        self.assertEqual(updated.status_code, 200)
        self.assertEqual(updated.data["design_options"], {"theme": "maroon", "font": "serif", "layout": "sidebar"})
        self.assertEqual(updated.data["template_key"], "professional-v1")

    def test_foreign_user_cannot_read_preview_or_download(self):
        job = ExportJob.objects.create(owner=self.user, document=self.document, purpose=ExportJob.PREVIEW,
            document_revision=self.document.revision, content_hash="a" * 64, render_input={}, status=ExportJob.SUCCEEDED)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f"/api/v1/exports/{job.id}/preview/").status_code, 404)
        self.assertEqual(self.client.get(f"/api/v1/exports/{job.id}/download/").status_code, 404)

    def test_export_download_requires_owner_and_succeeded_export(self):
        job = ExportJob.objects.create(owner=self.user, document=self.document, purpose=ExportJob.EXPORT,
            document_revision=self.document.revision, content_hash="b" * 64, render_input={}, status=ExportJob.SUCCEEDED)
        job.artifact.save("test.pdf", ContentFile(b"%PDF-1.4 test"), save=True)
        self.assertEqual(self.client.get(f"/api/v1/exports/{job.id}/download/").status_code, 200)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(f"/api/v1/exports/{job.id}/download/").status_code, 404)

    def test_server_usage_reports_active_charges_and_expires_old_events(self):
        from datetime import timedelta
        from django.utils import timezone
        from modules.resumes.models import ExportCharge
        from .quota import usage_for
        now = timezone.now()
        ExportCharge.objects.create(owner=self.user, document_identity=uuid.uuid4(), expires_at=now + timedelta(hours=4))
        ExportCharge.objects.create(owner=self.user, document_identity=uuid.uuid4(), expires_at=now - timedelta(seconds=1))
        usage = usage_for(self.user, now)
        self.assertEqual(usage["plan"].code, "free")
        self.assertEqual(usage["used"], 1)
        self.assertEqual(usage["reserved"], 0)
        self.assertEqual(usage["available"], 4)
        self.client.get("/api/v1/exports/usage/")
        self.assertEqual(self.client.get("/api/v1/exports/usage/").data["used"], 1)

    @patch("modules.exports.views.render_export.delay")
    def test_new_export_reserves_one_credit_then_settles_only_on_success(self, delay):
        from .quota import usage_for
        from unittest.mock import patch as mock_patch
        from modules.resumes.models import ExportCharge
        self.document.snapshot = {"personal": {"full_name": "Nasrat"}}
        self.document.save(update_fields=["snapshot"])
        key = str(uuid.uuid4())
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post("/api/v1/exports/", {"document_id": str(self.document.id), "purpose": "export", "idempotency_key": key}, format="json")
        self.assertEqual(response.status_code, 202)
        job = ExportJob.objects.get(pk=response.data["id"])
        self.assertEqual(usage_for(self.user)["reserved"], 1)
        with mock_patch("modules.exports.tasks.render_cv_pdf", return_value=b"%PDF-1.4 acceptance"):
            from .tasks import render_export
            render_export.run(str(job.pk))
        self.assertEqual(ExportCharge.objects.filter(owner=self.user, document_identity=self.document.id).count(), 1)
        self.assertEqual(usage_for(self.user)["reserved"], 0)
        repeated = self.client.post("/api/v1/exports/", {"document_id": str(self.document.id), "purpose": "export", "idempotency_key": str(uuid.uuid4())}, format="json")
        self.assertEqual(repeated.data["id"], str(job.id))
        self.assertEqual(ExportCharge.objects.filter(owner=self.user, document_identity=self.document.id).count(), 1)

    def test_new_distinct_cv_is_denied_when_rolling_limit_is_full(self):
        from .quota import usage_for
        from datetime import timedelta
        from django.utils import timezone
        from modules.resumes.models import ExportCharge
        for _ in range(5):
            ExportCharge.objects.create(owner=self.user, document_identity=uuid.uuid4(), expires_at=timezone.now()+timedelta(hours=12))
        response = self.client.post("/api/v1/exports/", {"document_id": str(self.document.id), "purpose": "export", "idempotency_key": str(uuid.uuid4())}, format="json")
        self.assertEqual(response.status_code, 429)
        self.assertEqual(response.data["code"], "quota_exhausted")
        self.assertEqual(usage_for(self.user)["available"], 0)

    def test_unverified_or_expired_premium_subscription_resolves_to_free(self):
        from datetime import timedelta
        from django.utils import timezone
        from modules.billing.models import Plan, Subscription
        from modules.billing.services import effective_entitlement
        now = timezone.now()
        premium = Plan.objects.get(code="premium")
        subscription = Subscription.objects.create(owner=self.user, plan=premium, starts_at=now-timedelta(days=1), expires_at=now+timedelta(days=10))
        self.assertEqual(effective_entitlement(self.user)["plan"].code, "free")
        subscription.verified_at = now
        subscription.expires_at = now-timedelta(seconds=1)
        subscription.save(update_fields=["verified_at", "expires_at"])
        self.assertEqual(effective_entitlement(self.user)["plan"].code, "free")
        subscription.expires_at = now+timedelta(days=10)
        subscription.save(update_fields=["expires_at"])
        self.assertEqual(effective_entitlement(self.user)["plan"].code, "premium")
