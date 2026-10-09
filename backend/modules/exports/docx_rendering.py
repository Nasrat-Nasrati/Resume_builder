"""Editable DOCX output from the immutable export contract."""
import base64
import binascii
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .templates import LABELS, THEMES, template_for, validate_design_options, default_design_options


def _direction(paragraph, rtl):
    if rtl:
        p_pr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        p_pr.append(bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _run(paragraph, text, *, bold=False, size=None, color=None, rtl=False, font="sans"):
    run = paragraph.add_run(str(text))
    run.bold = bold
    family = ("Noto Naskh Arabic" if rtl else "Noto Serif") if font == "serif" else ("Noto Naskh Arabic" if rtl else "Noto Sans")
    run.font.name = family
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    r_pr = run._r.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    fonts.set(qn("w:ascii"), family)
    fonts.set(qn("w:hAnsi"), family)
    fonts.set(qn("w:cs"), family)
    if rtl:
        rtl_tag = OxmlElement("w:rtl")
        rtl_tag.set(qn("w:val"), "1")
        r_pr.append(rtl_tag)
    return run


def _value(value):
    if value is None or value == "" or isinstance(value, bool):
        return ""
    if isinstance(value, list):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    return str(value)


def render_cv_docx(payload):
    key = payload.get("template_key", "public-service-v1")
    template = template_for(key)
    language = payload.get("output_language", "en")
    if language not in template["supported_languages"]:
        raise ValueError("Unsupported output language.")
    design = validate_design_options(key, payload.get("design_options", default_design_options(key)))
    rtl = language in {"fa", "ps"}
    labels = LABELS[language]
    snapshot = payload.get("snapshot", {})
    personal = snapshot.get("personal", {})
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    colors = THEMES[design["theme"]]
    ink, accent = colors["ink"], colors["accent"]
    font = design["font"]
    normal = document.styles["Normal"]
    normal.font.name = "Noto Naskh Arabic" if rtl else ("Noto Serif" if font == "serif" else "Noto Sans")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(ink.lstrip("#"))

    p = document.add_paragraph()
    _direction(p, rtl)
    _run(p, personal.get("full_name") or payload.get("title") or "CV", bold=True, size=24, color=ink, rtl=rtl, font=font)
    photo_data = payload.get("photo_data") if key != "ats-v1" else None
    if photo_data:
        if not isinstance(photo_data, str) or len(photo_data) > 700_000:
            raise ValueError("Invalid embedded photo data.")
        try:
            photo_bytes = base64.b64decode(photo_data, validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ValueError("Invalid embedded photo data.") from exc
        if len(photo_bytes) > 512 * 1024:
            raise ValueError("Embedded photo exceeds the supported size.")
        photo_paragraph = document.add_paragraph()
        photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
        photo_paragraph.add_run().add_picture(BytesIO(photo_bytes), width=Inches(1.2))
    title = personal.get("professional_title")
    if title:
        p = document.add_paragraph()
        _direction(p, rtl)
        _run(p, title, size=12, color=accent, rtl=rtl, font=font)
    contact = " · ".join(str(personal.get(field)) for field in ("email", "phone") if personal.get(field))
    if contact:
        p = document.add_paragraph()
        _direction(p, rtl)
        _run(p, contact, size=9, font=font)
    location = " · ".join(str(personal.get(field)) for field in ("city", "province", "district") if personal.get(field) and (field == "city" or personal.get(f"show_{field}", True)))
    if location:
        p = document.add_paragraph()
        _direction(p, rtl)
        _run(p, location, size=9, rtl=rtl, font=font)
    afghan = " · ".join(str(personal.get(field)) for field in ("father_name", "tazkira_number") if personal.get(field) and personal.get(f"show_{field}", False))
    if afghan:
        p = document.add_paragraph()
        _direction(p, rtl)
        _run(p, afghan, size=9, rtl=rtl, font=font)

    if design["layout"] == "sidebar":
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(2.0)
        table.columns[1].width = Inches(4.8)
        tbl_pr = table._tbl.tblPr
        bidi = OxmlElement("w:bidiVisual")
        bidi.set(qn("w:val"), "1" if rtl else "0")
        tbl_pr.append(bidi)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "nil")
            borders.append(element)
        tbl_pr.append(borders)
        sidebar, main = table.rows[0].cells
        target_for = lambda section_key: sidebar if section_key in {"profile", "languages", "skills"} else main
    else:
        sidebar = main = document
        target_for = lambda section_key: document

    def add_paragraph(container, text, *, bold=False, size=None, color=None, keep=False):
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.keep_with_next = keep
        _direction(paragraph, rtl)
        _run(paragraph, text, bold=bold, size=size, color=color, rtl=rtl, font=font)
        return paragraph

    def add_heading(container, text):
        paragraph = add_paragraph(container, text, bold=True, size=13, color=accent if key != "ats-v1" else "222222")
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(3)

    if personal.get("summary"):
        target = target_for("profile")
        add_heading(target, labels["profile"])
        add_paragraph(target, personal["summary"])
    selected = set(payload.get("selected_sections", []))
    for section_key in template["sections"][1:]:
        if section_key in {"projects", "awards", "publications", "references"} and section_key not in selected:
            continue
        rows = snapshot.get(section_key, [])
        if not rows:
            continue
        target = target_for(section_key)
        visible = [row for row in rows if not (section_key == "references" and isinstance(row, dict) and not row.get("consent_to_publish"))]
        if not visible:
            continue
        add_heading(target, labels[section_key])
        for row in visible:
            if not isinstance(row, dict):
                add_paragraph(target, row)
                continue
            fields = ("referee_name", "referee_role") if section_key == "references" else ("position", "qualification", "title", "name", "language", "institution", "employer", "organization")
            heading = " — ".join(str(row[field]) for field in fields if row.get(field))
            if heading:
                add_paragraph(target, heading, bold=True, keep=True)
            dates = " – ".join(str(row[field]) for field in ("start_date", "end_date", "date") if row.get(field))
            if dates:
                add_paragraph(target, dates, size=9, color="5D6B7B")
            detail = row.get("duties") or row.get("description") or row.get("field_of_study")
            if detail:
                add_paragraph(target, detail)
            if section_key == "skills" and row.get("skills"):
                add_paragraph(target, _value(row["skills"]))
            if section_key == "references":
                private_contact = " · ".join(str(row[field]) for field in ("referee_email", "referee_phone") if row.get(field))
                if private_contact:
                    p = target.add_paragraph()
                    _direction(p, rtl)
                    _run(p, private_contact, size=9, font=font)
    if payload.get("watermark"):
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(footer, "KabulCV · Free", size=8, color="708090", font=font)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
